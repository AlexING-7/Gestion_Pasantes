
import sys
from pathlib import Path
sys.path.append(str(Path(__file__).resolve().parent.parent))
from docxtpl import DocxTemplate,InlineImage
from docx.shared import Mm, Cm, Inches
import jinja2
from modelos.modulo import Pasantia,Enterprise,session,Configuracion
import datetime
import os
from herramientas.conversiones import mes_espanol,convertir_a_romano,fecha_espanol,formato_miles, nombreCompleto,nombreParcial,sede_Sucursal

def reemplazar_texto(datos:Pasantia,filename,name):
    conf=session.query(Configuracion).all()
    dic_conf=dict((i.clave,i.valor) for i in conf)
    doc = DocxTemplate(filename)
    jinja_env = jinja2.Environment()
    jinja_env.filters['fecha_es'] = fecha_espanol
    jinja_env.filters['miles'] = formato_miles
    jinja_env.filters['romanos'] = convertir_a_romano
    jinja_env.filters['mes'] = mes_espanol
    jinja_env.filters['completo'] = nombreCompleto
    jinja_env.filters['parcial'] = nombreParcial
    jinja_env.filters['direccion'] = sede_Sucursal
    perfil_pasante=InlineImage(doc,datos.student.foto,width=Mm(32),height=Mm(35))
    contexto={
        'hoy':datetime.datetime.now(),
        #datos
        'datos':datos,
        'estudiante':datos.student,
        'foto_estudiante':perfil_pasante,
        'empresa':datos.empresa,
        'tutor_academico':datos.tutor_academico,
        'tutor_empresarial':datos.tutor_empresarial,
        'carrera': datos.carrera,
        'semestre': datos.semestre,
        'lapso': datos.lapso_academico,
        'inicio': datos.inicio_pasantias,
        'final': datos.final_pasantias,
        'departamento': datos.departamento,
        'estado': datos.estado,
        'trabajo_asignado': datos.trabajo_asignado,
        'titulo_de_informe': datos.titulo_de_informe,
        'autoridad_empresa':datos.jefe_de_carta,
        'cargo_empresa':datos.cargo_jefe_de_carta
       }
    contexto = {k: (v if v is not None else "N/A") for k, v in contexto.items()}
    contexto.update(dic_conf)

    # 3. Renderizar (rellenar) el documento
    doc.render(contexto,jinja_env)

    # 4. Guardar el resultado final
    nombre=name
    doc.save(nombre)

    try:
        os.startfile(nombre)
    except AttributeError:
        # Esto ocurre si intentas correr este código en Mac o Linux
        print("os.startfile solo funciona en Windows.")

def asignacion_pasantias(datos,name="prueba.docx"):
    conf=session.query(Configuracion).all()
    dic_conf=dict((i.clave,i.valor) for i in conf)
    jinja_env = jinja2.Environment()
    jinja_env.filters['fecha_es'] = fecha_espanol
    jinja_env.filters['miles'] = formato_miles
    jinja_env.filters['romanos'] = convertir_a_romano
    jinja_env.filters['mes'] = mes_espanol
    jinja_env.filters['completo'] = nombreCompleto
    jinja_env.filters['parcial'] = nombreParcial
    jinja_env.filters['direccion'] = sede_Sucursal
    doc = DocxTemplate("formatos/ASIGNACION DE PASANTES.docx")
    
    contexto={"datos":datos,'hoy':datetime.datetime.now()}
    contexto.update(dic_conf)
    doc.render(contexto,jinja_env)
    doc.save(name)
    
    try:
        os.startfile(name)
    except AttributeError:
        # Esto ocurre si intentas correr este código en Mac o Linux
        print("os.startfile solo funciona en Windows.")

from openpyxl import load_workbook

def evaluacionAcademico(datos:Pasantia,formato,filename,path):
    wb = load_workbook(filename)
    ws = wb.active  # O wb['Nombre de la hoja']
    conf=session.query(Configuracion).all()
    dic_conf=dict((i.clave,i.valor) for i in conf)
    # 2. Rellenar los campos de texto
    # Nota: Debes verificar en qué celda exacta cae cada campo en tu Excel
    if 'Evaluación Tutor Académico'==formato:
        ws['B9'] = f"{datos.student.primer_apellido} {getattr(datos.student,"segundo_apellido","")}"    
        ws['D9'] = f"{datos.student.primer_nombre} {getattr(datos.student,"segundo_nombre","")}"        
        ws['G9'] = f"{datos.student.cedula}"    
        ws['B10'] = f"{dic_conf["ciudad"]}"     
        ws['D10']=f"{datos.carrera}"
        ws['G10']=f"{datos.lapso_academico}"
        ws['B12'] = f"{datos.empresa.razon_social}"   
        ws['D12'] = f"{datos.inicio_pasantias}" 
        ws['G12'] = f"{datos.final_pasantias}" 
        ws['B13'] = f"{nombreCompleto(datos.tutor_empresarial)}" 
        ws['D13'] = f"{nombreCompleto(datos.tutor_academico)}" 

    if 'Evaluación Tutor Empresarial'==formato:
        ws['B5'] = f"{datos.student.primer_apellido} {getattr(datos.student,"segundo_apellido","")}"     
        ws['D5'] = f"{datos.student.primer_nombre} {getattr(datos.student,"segundo_nombre","")}"           
        ws['G5'] = f"{datos.student.cedula}"    
        ws['B6'] = f"{dic_conf["ciudad"]}"      
        ws['D6']=f"{datos.carrera}"
        ws['G6']=f"{datos.lapso_academico}"
        ws['B8'] = f"{datos.empresa.razon_social}"    
        ws['D8'] = f"{datos.inicio_pasantias}" 
        ws['G8'] = f"{datos.final_pasantias}" 
        ws['B9'] = f"{nombreCompleto(datos.tutor_academico)}" 
        ws['D9'] = f"{nombreCompleto(datos.tutor_empresarial)}" 
        ws['B10'] = f"{datos.tutor_empresarial.telefono}" 
        ws['D10'] = f"{datos.tutor_empresarial.email}" 


    # 4. Guardar como un archivo nuevo para no borrar la plantilla
    wb.save(path)
    os.startfile(path)

    